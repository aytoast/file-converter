"""Table data exchange with explicit worksheet and formula handling."""
import csv
import datetime
import html
import json
from pathlib import Path

INPUTS = {".csv", ".tsv", ".xlsx", ".xlsm", ".xls", ".ods", ".json", ".jsonl"}
TARGETS = {"csv", "tsv", "xlsx", "ods", "json", "jsonl", "html"}


def text(source, encoding):
    return source.read_text(encoding=encoding, errors="strict")


def read_tables(source, args):
    suffix = source.suffix.lower()
    warnings = []
    if suffix in {".csv", ".tsv"}:
        content = text(source, args.encoding)
        delimiter = args.delimiter or ("\t" if suffix == ".tsv" else ",")
        if len(delimiter) != 1:
            raise ValueError("Delimiter must be one character")
        import io
        rows = list(csv.reader(io.StringIO(content, newline=""), delimiter=delimiter, strict=True))
        return [(source.stem, rows)], warnings
    if suffix in {".json", ".jsonl"}:
        raw = text(source, args.encoding)
        records = (json.loads(raw) if suffix == ".json" else
                   [json.loads(line) for line in raw.splitlines() if line.strip()])
        if not isinstance(records, list) or not all(isinstance(row, dict) for row in records):
            raise ValueError("Tabular JSON input must be a list of flat objects")
        headers = list(dict.fromkeys(key for row in records for key in row))
        if any(isinstance(value, (dict, list)) for row in records for value in row.values()):
            raise ValueError("Nested JSON needs an explicit flattening schema")
        rows = [headers] + [[row.get(key) for key in headers] for row in records]
        return [(source.stem, rows)], warnings
    if suffix in {".xlsx", ".xlsm"}:
        import openpyxl
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
        try:
            names = workbook.sheetnames
            if len(names) > 1 and not (args.sheet or args.all_sheets):
                raise ValueError("Multiple worksheets: choose --sheet NAME or --all-sheets")
            if args.sheet and args.sheet not in names:
                raise ValueError("Requested worksheet does not exist")
            names = [args.sheet] if args.sheet else names
            tables = []
            formulas = False
            for name in names:
                sheet = workbook[name]
                rows = []
                for cells in sheet.iter_rows():
                    row = []
                    for cell in cells:
                        if cell.data_type == "f":
                            formulas = True
                            if args.formulas == "reject":
                                raise ValueError("Formulas present: choose --formulas values or expressions")
                        row.append(cell.value)
                    rows.append(row)
                tables.append((name, rows))
            if formulas and args.formulas == "values":
                cached = openpyxl.load_workbook(source, read_only=True, data_only=True)
                try:
                    for index, (name, rows) in enumerate(tables):
                        for row_index, row in enumerate(rows):
                            for col_index, value in enumerate(row):
                                cell = workbook[name].cell(row_index + 1, col_index + 1)
                                if cell.data_type == "f":
                                    value = cached[name].cell(row_index + 1, col_index + 1).value
                                    if value is None:
                                        raise ValueError("Formula has no cached value; recalculate in Excel first")
                                    rows[row_index][col_index] = value
                finally:
                    cached.close()
                warnings.append("Formula results are cached values, not freshly recalculated")
            if formulas and args.formulas == "expressions":
                warnings.append("Formula expressions exported as text; they will not recalculate")
            warnings.append("Data exchange only: workbook styling, charts and macros are not retained")
            return tables, warnings
        finally:
            workbook.close()
    if suffix == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(source)
        try:
            names = workbook.sheet_names()
            if len(names) > 1 and not (args.sheet or args.all_sheets):
                raise ValueError("Multiple worksheets: choose --sheet NAME or --all-sheets")
            selected = [args.sheet] if args.sheet else names
            tables = []
            for name in selected:
                sheet = workbook.sheet_by_name(name)
                rows = []
                for cells in sheet.get_rows():
                    row = []
                    for cell in cells:
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate.xldate_as_datetime(value, workbook.datemode)
                        row.append(value)
                    rows.append(row)
                tables.append((name, rows))
            return tables, ["Legacy XLS data only; formulas are read as stored results"]
        finally:
            workbook.release_resources()
    if suffix == ".ods":
        from odf.opendocument import load
        from odf import table, teletype
        document = load(str(source))
        sheets = document.spreadsheet.getElementsByType(table.Table)
        names = [s.getAttribute("name") for s in sheets]
        if len(sheets) > 1 and not (args.sheet or args.all_sheets):
            raise ValueError("Multiple worksheets: choose --sheet NAME or --all-sheets")
        if args.sheet and args.sheet not in names:
            raise ValueError("Requested worksheet does not exist")
        result = []
        for sheet in sheets:
            if args.sheet and sheet.getAttribute("name") != args.sheet:
                continue
            rows = []
            for row in sheet.getElementsByType(table.TableRow):
                values = []
                for cell in row.childNodes:
                    if getattr(cell, "qname", None) not in {table.TableCell().qname, table.CoveredTableCell().qname}:
                        continue
                    if cell.getAttribute("formula"):
                        if args.formulas == "reject":
                            raise ValueError("ODS formulas require --formulas values or expressions")
                        if args.formulas == "expressions":
                            value = cell.getAttribute("formula")
                        else:
                            value = cell.getAttribute("value") or cell.getAttribute("stringvalue") or teletype.extractText(cell)
                    else:
                        value = cell.getAttribute("value") or cell.getAttribute("stringvalue") or teletype.extractText(cell)
                    repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
                    if repeat > 16384 or len(values) + repeat > 16384:
                        raise ValueError("ODS repeated columns exceed supported data bounds")
                    values.extend([value] * repeat)
                repeat_rows = int(row.getAttribute("numberrowsrepeated") or 1)
                if repeat_rows > 100000 or len(rows) + repeat_rows > 100000:
                    raise ValueError("ODS repeated rows exceed supported data bounds")
                rows.extend([list(values) for _ in range(repeat_rows)])
            result.append((sheet.getAttribute("name"), rows))
        return result, ["ODS data only; formatting and formula recalculation are not retained"]
    raise ValueError("Unsupported table input")


def scalar(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    return value


def records(rows):
    if not rows:
        return []
    headers = [str(value) if value is not None else "" for value in rows[0]]
    if any(not h for h in headers) or len(set(headers)) != len(headers):
        raise ValueError("JSON export needs unique nonempty column headers")
    if any(len(row) > len(headers) for row in rows[1:]):
        raise ValueError("Row has more values than headers")
    return [{key: scalar(row[index]) if index < len(row) else None
             for index, key in enumerate(headers)} for row in rows[1:]]


def write_tables(tables, workspace, target, args):
    outputs = []
    if target == "xlsx":
        from openpyxl import Workbook, load_workbook
        book = Workbook()
        book.remove(book.active)
        for name, rows in tables:
            sheet = book.create_sheet(name[:31] or "Sheet")
            if len(rows) > 1048576 or any(len(row) > 16384 for row in rows):
                raise ValueError("Data exceeds XLSX worksheet limits")
            for row in rows:
                if any(isinstance(value, str) and len(value) > 32767 for value in row):
                    raise ValueError("Cell text exceeds XLSX limit")
                row = [str(value) if isinstance(value, int) and not isinstance(value, bool) and abs(value) >= 10**15
                       else value for value in row]
                sheet.append(row)
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
        destination = workspace / "converted.xlsx"
        book.save(destination)
        check = load_workbook(destination, read_only=True)
        try:
            if len(check.sheetnames) != len(tables):
                raise ValueError("XLSX worksheet verification failed")
        finally:
            check.close()
        return [destination]
    if target == "ods":
        from odf.opendocument import OpenDocumentSpreadsheet, load
        from odf.table import Table, TableRow, TableCell
        from odf.text import P
        book = OpenDocumentSpreadsheet()
        for name, rows in tables:
            sheet = Table(name=name)
            for values in rows:
                row = TableRow()
                for value in values:
                    cell = TableCell(valuetype="string")
                    cell.addElement(P(text="" if value is None else str(scalar(value))))
                    row.addElement(cell)
                sheet.addElement(row)
            book.spreadsheet.addElement(sheet)
        destination = workspace / "converted.ods"
        book.save(str(destination))
        load(str(destination))
        return [destination]
    for index, (name, rows) in enumerate(tables):
        destination = workspace / f"sheet-{index+1:03d}.{target}"
        if target in {"csv", "tsv"}:
            with destination.open("w", encoding=args.output_encoding, newline="") as output:
                writer = csv.writer(output, delimiter="\t" if target == "tsv" else (args.delimiter or ","))
                writer.writerows([[scalar(v) for v in row] for row in rows])
        elif target in {"json", "jsonl"}:
            content = records(rows)
            encoded = (json.dumps(content, ensure_ascii=False, allow_nan=False, indent=2) if target == "json" else
                       "\n".join(json.dumps(record, ensure_ascii=False, allow_nan=False) for record in content))
            destination.write_text(encoded, encoding="utf-8")
        elif target == "html":
            body = "".join("<tr>" + "".join("<td>" + html.escape("" if value is None else str(scalar(value))) +
                                            "</td>" for value in row) + "</tr>" for row in rows)
            destination.write_text('<!doctype html><meta charset="utf-8"><table>' + body + "</table>", encoding="utf-8")
        outputs.append(destination)
    return outputs


def convert(source, workspace, target, args):
    tables, warnings = read_tables(source, args)
    if not tables:
        raise ValueError("No tables found")
    paths = write_tables(tables, workspace, target, args)
    return paths, warnings, "table_output_reopened"
