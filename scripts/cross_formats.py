"""Explicit cross-category extraction, OCR, transcription and image sequencing."""
import functools
import json
import os
from pathlib import Path
import zipfile


@functools.lru_cache(maxsize=1)
def ocr_engine():
    from rapidocr_onnxruntime import RapidOCR
    return RapidOCR(intra_op_num_threads=2, inter_op_num_threads=2,
                    det_use_cuda=False, cls_use_cuda=False, rec_use_cuda=False)


def recognize(image):
    import numpy as np
    # RapidOCR accepts BGR arrays; normalize RGB/RGBA with an explicit white page.
    from PIL import Image
    rgba = image.convert("RGBA")
    background = Image.new("RGB", rgba.size, "white")
    background.paste(rgba, mask=rgba.getchannel("A"))
    result, _ = ocr_engine()(np.asarray(background)[:, :, ::-1].copy())
    return [{"box": [[float(x), float(y)] for x, y in box], "text": str(text), "score": float(score)}
            for box, text, score in (result or [])]


def ocr_pages(source, args):
    from PIL import Image, ImageOps, ImageSequence
    pages = []
    if source.suffix.lower() == ".pdf":
        import pymupdf
        from document_formats import selected_pages
        with pymupdf.open(source) as document:
            if document.needs_pass:
                raise ValueError("PDF requires a password")
            indices = selected_pages(document, args.pages)
            for index in indices:
                pixmap = document[index].get_pixmap(matrix=pymupdf.Matrix(args.dpi / 72, args.dpi / 72), alpha=False)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                pages.append({"index": index, "size": image.size, "records": recognize(image)})
    else:
        with Image.open(source) as original:
            if getattr(original, "n_frames", 1) > 1 and not args.all_frames:
                raise ValueError("Multi-page OCR requires --all-frames")
            for index, frame in enumerate(ImageSequence.Iterator(original)):
                image = ImageOps.exif_transpose(frame).convert("RGB")
                pages.append({"index": index, "size": image.size, "records": recognize(image)})
    if not any(page["records"] for page in pages):
        raise ValueError("OCR detected no text; no editable output was created")
    return pages


def table_rows(records):
    if not records:
        return []
    import statistics
    heights = [max(p[1] for p in item["box"]) - min(p[1] for p in item["box"]) for item in records]
    tolerance = max(5, statistics.median(heights) * 0.65)
    rows = []
    for item in sorted(records, key=lambda r: min(p[1] for p in r["box"])):
        center = sum(p[1] for p in item["box"]) / 4
        row = next((row for row in rows if abs(row[0] - center) <= tolerance), None)
        if row is None:
            row = [center, []]
            rows.append(row)
        row[1].append(item)
    starts = []
    for _, values in rows:
        for item in values:
            x = min(point[0] for point in item["box"])
            if not any(abs(existing - x) < tolerance * 2 for existing in starts):
                starts.append(x)
    starts.sort()
    if len(starts) < 2 or len(rows) < 2:
        raise ValueError("OCR found no reliable multi-column table; choose text output instead")
    output = []
    for _, values in rows:
        cells = [""] * len(starts)
        for item in sorted(values, key=lambda r: min(p[0] for p in r["box"])):
            x = min(point[0] for point in item["box"])
            column = min(range(len(starts)), key=lambda index: abs(starts[index] - x))
            cells[column] = (cells[column] + " " + item["text"]).strip()
        output.append(cells)
    return output


def ocr(source, workspace, target, args):
    pages = ocr_pages(source, args)
    warnings = ["OCR is inferred text; verify names, numbers, reading order and low-confidence recognition"]
    low = sum(item["score"] < 0.8 for page in pages for item in page["records"])
    if low:
        warnings.append(f"{low} recognized text regions have confidence below 0.8")
    output = workspace / f"converted.{target}"
    if target in {"txt", "md"}:
        output.write_text("\n\n\f\n\n".join("\n".join(item["text"] for item in page["records"])
                                                    for page in pages), encoding="utf-8")
    elif target == "json":
        output.write_text(json.dumps(pages, ensure_ascii=False, indent=2), encoding="utf-8")
    elif target == "docx":
        from docx import Document
        document = Document()
        for number, page in enumerate(pages):
            if number:
                document.add_page_break()
            for item in page["records"]:
                document.add_paragraph(item["text"])
        document.save(output)
        Document(output)
        warnings.append("Editable recognized text; original page layout is not reconstructed")
    elif target in {"csv", "tsv", "xlsx"}:
        from table_formats import write_tables
        tables = [(f"Page{page['index']+1}", table_rows(page["records"])) for page in pages if page["records"]]
        outputs = write_tables(tables, workspace, target, args)
        return outputs, warnings + ["Table rows/columns inferred from text positions; merged cells are not restored"], "ocr_table_output_reopened"
    elif target == "pdf":
        import pymupdf
        from PIL import Image, ImageOps, ImageSequence
        document = pymupdf.open()
        if source.suffix.lower() == ".pdf":
            with pymupdf.open(source) as original:
                for page in pages:
                    document.insert_pdf(original, from_page=page["index"], to_page=page["index"])
        else:
            import io
            with Image.open(source) as original:
                for frame in ImageSequence.Iterator(original):
                    image = ImageOps.exif_transpose(frame).convert("RGB")
                    data = io.BytesIO()
                    image.save(data, format="PNG")
                    page = document.new_page(width=image.width * 72 / args.dpi, height=image.height * 72 / args.dpi)
                    page.insert_image(page.rect, stream=data.getvalue())
        for index, content in enumerate(pages):
            page = document[index]
            # Existing text remains intact; don't duplicate native text layers.
            if page.get_text().strip():
                continue
            sx, sy = page.rect.width / content["size"][0], page.rect.height / content["size"][1]
            for item in content["records"]:
                xs = [p[0] * sx for p in item["box"]]
                ys = [p[1] * sy for p in item["box"]]
                width, height = max(xs) - min(xs), max(ys) - min(ys)
                units = pymupdf.get_text_length(item["text"], fontname="china-s", fontsize=1)
                size = max(0.5, min(height * 0.85, width / max(units, 0.1)))
                page.insert_text((min(xs), max(ys) - height * 0.15), item["text"],
                                 fontname="china-s", fontsize=size, render_mode=3)
        document.save(output)
        document.close()
        with pymupdf.open(output) as check:
            if not any(page.get_text().strip() for page in check):
                raise ValueError("Searchable PDF text-layer verification failed")
    else:
        raise ValueError("Unsupported OCR target")
    return [output], warnings, "ocr_text_output_checked"


@functools.lru_cache(maxsize=2)
def asr_engine(directory):
    from faster_whisper import WhisperModel
    return WhisperModel(directory, device="cpu", compute_type="int8", cpu_threads=2, local_files_only=True)


def asr(source, workspace, target, args):
    configured = args.asr_model_path or os.environ.get("FILE_CONVERTER_ASR_MODEL")
    if not configured or not Path(configured).is_dir():
        raise ValueError("ASR requires a local model directory via --asr-model-path; conversion never downloads models")
    model = asr_engine(str(Path(configured).resolve()))
    if args.language and args.language != "en" and not model.model.is_multilingual:
        raise ValueError("Selected ASR model is English-only; choose a multilingual model for this language")
    segments, info = model.transcribe(str(source), language=args.language, beam_size=3, vad_filter=False)
    records = [{"start": segment.start, "end": segment.end, "text": segment.text.strip()} for segment in segments]
    records = [item for item in records if item["text"]]
    if not records:
        raise ValueError("ASR produced no speech text")
    output = workspace / f"converted.{target}"
    if target in {"txt", "md"}:
        output.write_text("\n".join(item["text"] for item in records), encoding="utf-8")
    elif target == "json":
        output.write_text(json.dumps({"language": info.language, "segments": records}, ensure_ascii=False, indent=2), encoding="utf-8")
    elif target == "docx":
        from docx import Document
        document = Document()
        for item in records:
            document.add_paragraph(item["text"])
        document.save(output)
        Document(output)
    elif target in {"srt", "vtt"}:
        import pysubs2
        captions = pysubs2.SSAFile()
        for item in records:
            captions.append(pysubs2.SSAEvent(start=round(item["start"] * 1000),
                                            end=round(item["end"] * 1000), text=item["text"]))
        captions.save(str(output), encoding="utf-8", format_=target)
        if len(pysubs2.load(str(output), encoding="utf-8")) != len(records):
            raise ValueError("ASR subtitle count changed")
    else:
        raise ValueError("Unsupported ASR target")
    return [output], ["Speech recognition is inferred; verify wording, proper nouns and timing"], "asr_output_checked"


def extract_images(source, workspace):
    outputs = []
    if source.suffix.lower() == ".pdf":
        import pymupdf
        with pymupdf.open(source) as document:
            seen = set()
            for page in document:
                for info in page.get_images(full=True):
                    xref = info[0]
                    if xref in seen:
                        continue
                    seen.add(xref)
                    item = document.extract_image(xref)
                    output = workspace / f"image-{len(outputs)+1:04d}.{item['ext']}"
                    output.write_bytes(item["image"])
                    outputs.append(output)
    else:
        prefix = {".docx": "word/media/", ".pptx": "ppt/media/", ".xlsx": "xl/media/"}.get(source.suffix.lower())
        if not prefix:
            raise ValueError("Image extraction supports PDF, DOCX, PPTX and XLSX")
        from PIL import Image
        import io
        with zipfile.ZipFile(source) as archive:
            for name in archive.namelist():
                if name.startswith(prefix) and not name.endswith("/"):
                    data = archive.read(name)
                    try:
                        with Image.open(io.BytesIO(data)) as image:
                            image.verify()
                    except Exception:
                        continue
                    output = workspace / f"image-{len(outputs)+1:04d}{Path(name).suffix.lower()}"
                    output.write_bytes(data)
                    outputs.append(output)
    if not outputs:
        raise ValueError("No supported embedded images found")
    return outputs, ["Embedded image extraction, not page rendering; vector assets and masks may need separate handling"], "embedded_images_extracted"


def frames(source, workspace, args):
    from convert import ffmpeg_path, run, MEDIA_DEMUXERS
    if args.duration is None or args.duration <= 0:
        raise ValueError("Frame sequence needs explicit --duration")
    if args.duration * args.fps > 10000:
        raise ValueError("Frame request exceeds prototype limit of 10000 images")
    pattern = workspace / "frame-%05d.png"
    run([ffmpeg_path(), "-v", "error", "-y", "-nostdin", "-protocol_whitelist", "file,pipe",
         "-ss", str(args.start or 0), "-f", MEDIA_DEMUXERS[source.suffix.lower()], "-i", str(source),
         "-t", str(args.duration), "-vf", f"fps={args.fps}", "-an", str(pattern)])
    outputs = sorted(workspace.glob("frame-*.png"))
    if not outputs:
        raise ValueError("No video frames extracted")
    from PIL import Image
    for output in outputs:
        with Image.open(output) as image:
            image.load()
    return outputs, ["PNG frame sequence at selected sampling rate; no audio"], "video_frames_decoded"


def images_video(sources, workspace, args):
    from convert import ffmpeg_path, run
    from PIL import Image, ImageOps, ImageSequence
    normalized = []
    size = None
    for source in sources:
        with Image.open(source) as original:
            animated = getattr(original, "n_frames", 1) > 1
            for frame in ImageSequence.Iterator(original):
                rgba = ImageOps.exif_transpose(frame).convert("RGBA")
                if rgba.getchannel("A").getextrema()[0] < 255 and args.background is None:
                    raise ValueError("Image video needs an explicit background for transparency")
                if size is None:
                    size = (rgba.width + rgba.width % 2, rgba.height + rgba.height % 2)
                image = Image.new("RGB", size, args.background or "white")
                fitted = ImageOps.contain(rgba, size)
                image.paste(fitted, ((size[0] - fitted.width) // 2, (size[1] - fitted.height) // 2), fitted.getchannel("A"))
                output = workspace / f"input-{len(normalized)+1:05d}.png"
                image.save(output)
                duration = frame.info.get("duration", 100) / 1000 if animated else 1 / args.fps
                normalized.append((output, max(duration, 0.01)))
    if not normalized:
        raise ValueError("No image frames")
    manifest = workspace / "frames.ffconcat"
    lines = ["ffconcat version 1.0"]
    for image, duration in normalized:
        lines += [f"file '{image.name}'", f"duration {duration:.6f}"]
    lines.append(f"file '{normalized[-1][0].name}'")
    manifest.write_text("\n".join(lines) + "\n", encoding="utf-8")
    output = workspace / "converted.mp4"
    run([ffmpeg_path(), "-v", "error", "-y", "-nostdin", "-protocol_whitelist", "file,pipe",
         "-f", "concat", "-safe", "1", "-i", str(manifest), "-fps_mode", "vfr",
         "-an", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(output)])
    run([ffmpeg_path(), "-v", "error", "-xerror", "-i", str(output), "-f", "null", "-"])
    return [output], ["Image sequence rendered with explicit order; no audio; differing sizes letterboxed"], "image_video_decoded"


def raster_vector(source, workspace, args):
    if not args.allow_layout_loss:
        raise ValueError("Bitmap tracing produces monochrome paths, not photo-faithful vectors; choose --allow-layout-loss")
    import cv2
    import numpy as np
    from PIL import Image, ImageOps
    with Image.open(source) as original:
        if getattr(original, "n_frames", 1) != 1:
            raise ValueError("Vector tracing accepts a single image")
        image = ImageOps.exif_transpose(original).convert("RGBA")
        background = Image.new("RGB", image.size, "white")
        background.paste(image, mask=image.getchannel("A"))
        gray = cv2.cvtColor(np.asarray(background), cv2.COLOR_RGB2GRAY)
    _, binary = cv2.threshold(gray, args.vector_threshold, 255, cv2.THRESH_BINARY_INV)
    contours, _ = cv2.findContours(binary, cv2.RETR_TREE, cv2.CHAIN_APPROX_SIMPLE)
    paths = []
    for contour in contours:
        if len(contour) < 3 or cv2.contourArea(contour) < 1:
            continue
        points = contour[:, 0, :].tolist()
        paths.append("M " + " L ".join(f"{x},{y}" for x, y in points) + " Z")
    if not paths:
        raise ValueError("No foreground contours found")
    output = workspace / "converted.svg"
    width, height = image.size
    output.write_text(f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" '
                      f'viewBox="0 0 {width} {height}"><path fill="black" fill-rule="evenodd" '
                      f'd="{" ".join(paths)}"/></svg>', encoding="utf-8")
    from defusedxml import ElementTree
    ElementTree.parse(output)
    return [output], ["Monochrome contour tracing; color, texture and exact anti-aliasing are not retained"], "vector_paths_created"
