"""PDF reconstruction, content extraction and native Office conversion."""
import importlib.util
import os
from pathlib import Path
import shutil
import subprocess
import sys

OFFICE_INPUTS = {".doc", ".docx", ".odt", ".rtf", ".xls", ".xlsx", ".xlsm", ".ods", ".ppt", ".pptx", ".odp"}


def office_paths():
    roots = [Path(os.environ.get("ProgramFiles", "C:/Program Files")),
             Path(os.environ.get("ProgramFiles(x86)", "C:/Program Files (x86)"))]
    result = {}
    for name, executable in [("word", "WINWORD.EXE"), ("excel", "EXCEL.EXE"), ("powerpoint", "POWERPNT.EXE")]:
        result[name] = next((str(root / relative / executable) for root in roots
                             for relative in ["Microsoft Office/root/Office16", "Microsoft Office/Office16"]
                             if (root / relative / executable).is_file()), None)
    return result


def libreoffice_path():
    for name in ["soffice", "libreoffice"]:
        found = shutil.which(name)
        if found:
            return found
    for root in [os.environ.get("ProgramFiles", "C:/Program Files"),
                 os.environ.get("ProgramFiles(x86)", "C:/Program Files (x86)")]:
        found = Path(root) / "LibreOffice/program/soffice.exe"
        if found.is_file():
            return str(found)
    return None


def office_convert(source, workspace, target):
    suffix = source.suffix.lower()
    kind = ("excel" if suffix in {".xls", ".xlsx", ".xlsm", ".ods"} else
            "powerpoint" if suffix in {".ppt", ".pptx", ".odp"} else "word")
    allowed = {"word": {"doc", "docx", "odt", "rtf", "pdf", "html", "txt"},
               "excel": {"xls", "xlsx", "ods", "pdf"},
               "powerpoint": {"ppt", "pptx", "odp", "pdf"}}
    if target not in allowed[kind]:
        raise ValueError("Office target not supported for this document type")
    if target == "html":
        # Native HTML exports can create linked assets. Use explicit text extraction instead
        # of publishing an HTML file whose companion resources would be missing.
        intermediate = workspace / "docx-stage"
        intermediate.mkdir()
        paths, _, _ = office_convert(source, intermediate, "docx")
        return docx_text(paths[0], workspace, "html")
    output = workspace / f"converted.{target}"
    if office_paths()[kind] and importlib.util.find_spec("win32com"):
        process = subprocess.run([sys.executable, str(Path(__file__).with_name("office_worker.py")),
                                  "--input", str(source), "--output", str(output), "--kind", kind],
                                 capture_output=True, timeout=180)
        if process.returncode:
            detail = process.stderr.decode("utf-8", "replace")
            raise ValueError("Office export failed: " + detail[-1000:])
    elif libreoffice_path():
        profile = (workspace / "lo-profile").as_uri()
        process = subprocess.run([libreoffice_path(), "-env:UserInstallation=" + profile,
                                  "--headless", "--convert-to", target, "--outdir", str(workspace), str(source)],
                                 capture_output=True, timeout=180)
        generated = workspace / (source.stem + "." + target)
        if process.returncode or not generated.is_file():
            raise ValueError("LibreOffice export failed: " + process.stderr.decode("utf-8", "replace")[-1000:])
        generated.rename(output)
    else:
        raise ValueError("Office conversion needs the corresponding Microsoft Office app or LibreOffice")
    if not output.is_file() or output.stat().st_size == 0:
        raise ValueError("Office produced no output")
    verify(output)
    return [output], ["Native Office conversion; complex layouts, macros and unsupported features require review"], "office_output_reopened"


def office_images(source, workspace, target, args):
    intermediate = workspace / "pdf-stage"
    intermediate.mkdir()
    paths, warnings, _ = office_convert(source, intermediate, "pdf")
    outputs, rendering_warnings, validation = pdf_convert(paths[0], workspace, target, args)
    return outputs, warnings + rendering_warnings, validation


def verify(output):
    suffix = output.suffix.lower()
    if suffix == ".pdf":
        import pymupdf
        with pymupdf.open(output) as document:
            if not document.page_count:
                raise ValueError("Empty PDF")
            for page in document:
                page.get_pixmap(matrix=pymupdf.Matrix(0.3, 0.3))
    elif suffix == ".docx":
        from docx import Document
        Document(output)
    elif suffix == ".xlsx":
        import openpyxl
        workbook = openpyxl.load_workbook(output, read_only=True)
        workbook.close()
    elif suffix == ".pptx":
        from pptx import Presentation
        Presentation(output)
    elif suffix in {".odt", ".ods", ".odp"}:
        from odf.opendocument import load
        load(str(output))
    elif suffix in {".doc", ".xls", ".ppt"}:
        if output.read_bytes()[:8] != bytes.fromhex("D0CF11E0A1B11AE1"):
            raise ValueError("Invalid legacy Office container")


def selected_pages(document, expression):
    if not expression:
        return list(range(len(document)))
    values = []
    for piece in expression.split(","):
        bounds = piece.split("-")
        start = int(bounds[0])
        end = int(bounds[1]) if len(bounds) == 2 else start
        if len(bounds) > 2 or start < 1 or end < start or end > len(document):
            raise ValueError("Page range is out of bounds")
        values.extend(range(start - 1, end))
    return list(dict.fromkeys(values))


def pdf_convert(source, workspace, target, args):
    import pymupdf
    with pymupdf.open(source) as document:
        if document.needs_pass:
            raise ValueError("PDF requires a password")
        pages = selected_pages(document, args.pages)
        if target in {"png", "jpg", "tiff", "svg"}:
            outputs = []
            for index in pages:
                output = workspace / f"page-{index+1:04d}.{target}"
                if target == "svg":
                    output.write_text(document[index].get_svg_image(), encoding="utf-8")
                else:
                    pixmap = document[index].get_pixmap(matrix=pymupdf.Matrix(args.dpi / 72, args.dpi / 72), alpha=False)
                    pixmap.pil_save(str(output), format={"png": "PNG", "jpg": "JPEG", "tiff": "TIFF"}[target])
                    from PIL import Image
                    with Image.open(output) as check:
                        check.load()
                outputs.append(output)
            return outputs, ["Page rendering; not extraction of original embedded images"], "pdf_pages_rendered"
        if target in {"txt", "md", "html"}:
            if any(not document[index].get_text().strip() for index in pages):
                raise ValueError("A selected PDF page has no extractable text; OCR is required")
            content = "\n\n".join(document[index].get_text("html" if target == "html" else "text", sort=True)
                                    for index in pages)
            output = workspace / f"converted.{target}"
            output.write_text(content, encoding="utf-8")
            return [output], ["Reading-order text extraction; Markdown output has no reconstructed headings"], "pdf_text_extracted"
        if target in {"csv", "tsv", "xlsx", "json", "jsonl"}:
            from table_formats import write_tables
            tables = []
            for index in pages:
                for number, table in enumerate(document[index].find_tables().tables):
                    tables.append((f"Page{index+1}_Table{number+1}", table.extract()))
            if not tables:
                raise ValueError("No text-based PDF tables detected; scanned tables need OCR")
            return write_tables(tables, workspace, target, args), ["PDF table structure is inferred; verify numbers and merged cells"], "pdf_tables_extracted"
        if target == "pptx":
            if not args.allow_layout_loss:
                raise ValueError("PDF to PPTX creates image slides; use --allow-layout-loss after choosing this mode")
            from pptx import Presentation
            from pptx.util import Pt
            deck = Presentation()
            first = document[pages[0]].rect
            deck.slide_width, deck.slide_height = Pt(first.width), Pt(first.height)
            for index in pages:
                image = workspace / f"slide-{index+1}.png"
                document[index].get_pixmap(matrix=pymupdf.Matrix(args.dpi / 72, args.dpi / 72)).save(image)
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                ratio = min(deck.slide_width / document[index].rect.width, deck.slide_height / document[index].rect.height)
                width, height = int(document[index].rect.width * ratio), int(document[index].rect.height * ratio)
                slide.shapes.add_picture(str(image), int((deck.slide_width - width) / 2),
                                         int((deck.slide_height - height) / 2), width, height)
            output = workspace / "converted.pptx"
            deck.save(output)
            verify(output)
            return [output], ["Image slides: text and vector objects are not editable"], "image_slide_deck_reopened"
        if target == "docx":
            if any(not document[index].get_text().strip() for index in pages):
                raise ValueError("Scanned or empty PDF page: OCR is required before editable DOCX reconstruction")
            # Reconstruct exactly the selected pages, including noncontiguous ranges.
            subset = pymupdf.open()
            for index in pages:
                subset.insert_pdf(document, from_page=index, to_page=index)
            selected = workspace / "selected.pdf"
            subset.save(selected)
            subset.close()
    if target == "docx":
        import contextlib
        import io
        from pdf2docx import Converter
        output = workspace / "converted.docx"
        with contextlib.redirect_stdout(io.StringIO()):
            converter = Converter(str(selected))
            try:
                converter.convert(str(output), multi_processing=False)
            finally:
                converter.close()
        verify(output)
        from docx import Document
        result = Document(output)
        if not any(p.text.strip() for p in result.paragraphs) and not result.tables:
            raise ValueError("PDF reconstruction produced no editable text")
        return [output], ["Editable layout reconstruction; complex layout and font fidelity need visual review"], "editable_docx_reopened"
    raise ValueError("Unsupported PDF target")


def docx_text(source, workspace, target):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    import html
    document = Document(source)
    lines = []
    for element in document.element.body:
        if element.tag.endswith("}p"):
            paragraph = Paragraph(element, document)
            value = paragraph.text
            if target == "md" and paragraph.style.name.startswith("Heading "):
                level = paragraph.style.name.split()[-1]
                if level.isdigit():
                    value = "#" * min(int(level), 6) + " " + value
            lines.append(html.escape(value) if target == "html" else value)
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            for row in table.rows:
                value = "\t".join(cell.text for cell in row.cells)
                lines.append(html.escape(value) if target == "html" else value)
    output = workspace / f"converted.{target}"
    content = "\n\n".join(lines)
    if target == "html":
        content = '<!doctype html><meta charset="utf-8"><pre>' + content + "</pre>"
    output.write_text(content, encoding="utf-8")
    return [output], ["Text extraction: images, layout, footnotes and embedded objects not retained"], "docx_text_extracted"
