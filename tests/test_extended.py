import argparse
import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "convert.py"
spec = importlib.util.spec_from_file_location("extended_converter", SCRIPT)
converter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(converter)


class ExtendedTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.available = converter.capabilities()

    def setUp(self):
        self.workspace = tempfile.TemporaryDirectory(prefix="converter-expanded-")
        self.root = Path(self.workspace.name)

    def tearDown(self):
        self.assertEqual(list(self.root.glob(".file-converter-*")), [])
        self.workspace.cleanup()

    def convert(self, source, target, **options):
        before = hashlib.sha256(source.read_bytes()).hexdigest()
        args = argparse.Namespace(format=target, **options)
        try:
            result = converter.convert(source, args, self.available)
            self.assertEqual(result["status"], "success")
            return result
        finally:
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)

    def test_csv_xlsx_preserves_ids_and_literal_formulas(self):
        import openpyxl
        source = self.root / "表格 & data.csv"
        source.write_text('id,value\n0000123,12345678901234567890\n=1+1,"a,b"\n', encoding="utf-8-sig")
        result = self.convert(source, "xlsx")
        workbook = openpyxl.load_workbook(result["output"])
        self.assertEqual(workbook.active["A2"].value, "0000123")
        self.assertEqual(workbook.active["B2"].value, "12345678901234567890")
        self.assertEqual(workbook.active["A3"].data_type, "s")
        workbook.close()
        result = self.convert(Path(result["output"]), "csv")
        self.assertIn("0000123", Path(result["output"]).read_text(encoding="utf-8-sig"))

    def test_multiple_sheets_require_selection(self):
        import openpyxl
        source = self.root / "many.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["first"])
        workbook.create_sheet("Second").append(["second"])
        workbook.save(source)
        with self.assertRaisesRegex(ValueError, "Multiple worksheets"):
            self.convert(source, "csv")
        result = self.convert(source, "csv", all_sheets=True)
        self.assertEqual(len(result["outputs"]), 2)
        self.assertIn("second", Path(result["outputs"][1]).read_text(encoding="utf-8-sig"))

    def test_xlsx_limits_reject_truncation_and_preserve_long_json_integers(self):
        import openpyxl
        source = self.root / "long.json"
        source.write_text(json.dumps([{"id": 12345678901234567890}]), encoding="utf-8")
        result = self.convert(source, "xlsx")
        workbook = openpyxl.load_workbook(result["output"])
        self.assertEqual(workbook.active["A2"].value, "12345678901234567890")
        workbook.close()
        source.write_text(json.dumps([{"value": "x" * 32768}]), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "exceeds XLSX"):
            self.convert(source, "xlsx")

    def test_cli_pdf_json_and_target_discovery(self):
        source = self.pdf()
        result = subprocess.run([sys.executable, str(SCRIPT), "--input", str(source), "--format", "docx"],
                                capture_output=True)
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8", "replace"))
        self.assertEqual(json.loads(result.stdout.decode("utf-8"))["results"][0]["status"], "success")
        result = subprocess.run([sys.executable, str(SCRIPT), "--input", str(source), "--targets"], capture_output=True)
        self.assertEqual(result.returncode, 0)
        targets = json.loads(result.stdout.decode("utf-8"))["files"][0]["targets"]
        self.assertIn("docx", targets)
        self.assertIn("xlsx", targets)

    def test_uncalculated_formula_is_rejected(self):
        import openpyxl
        source = self.root / "formulas.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["value"])
        workbook.active.append(["=1+1"])
        workbook.save(source)
        with self.assertRaisesRegex(ValueError, "Formulas present"):
            self.convert(source, "csv")
        with self.assertRaisesRegex(ValueError, "no cached"):
            self.convert(source, "csv", formulas="values")
        result = self.convert(source, "csv", formulas="expressions")
        self.assertIn("=1+1", Path(result["output"]).read_text(encoding="utf-8-sig"))

    def test_json_table_and_ods_roundtrip(self):
        import openpyxl
        source = self.root / "records.json"
        source.write_text('[{"code":"0001","name":"中文"}]', encoding="utf-8")
        xlsx = Path(self.convert(source, "xlsx")["output"])
        result = self.convert(xlsx, "json")
        self.assertEqual(json.loads(Path(result["output"]).read_text(encoding="utf-8"))[0]["code"], "0001")
        csv = self.root / "records.csv"
        csv.write_text("code,name\n0001,中文\n", encoding="utf-8")
        ods = Path(self.convert(csv, "ods")["output"])
        converted = Path(self.convert(ods, "xlsx")["output"])
        workbook = openpyxl.load_workbook(converted)
        self.assertEqual(workbook.active["A2"].value, "0001")
        workbook.close()

    def test_json_yaml_roundtrip_and_duplicate_keys(self):
        source = self.root / "settings.json"
        source.write_text('{"name":"中文","items":[1,true,null]}', encoding="utf-8")
        yaml = Path(self.convert(source, "yaml")["output"])
        restored = Path(self.convert(yaml, "json")["output"])
        self.assertEqual(json.loads(source.read_text(encoding="utf-8")), json.loads(restored.read_text(encoding="utf-8")))
        source.write_text('{"a":1,"a":2}', encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Duplicate"):
            self.convert(source, "yaml")

    def pdf(self):
        import pymupdf
        source = self.root / "text.pdf"
        document = pymupdf.open()
        for content in ["First page editable text", "Second page selectable text"]:
            page = document.new_page()
            page.insert_text((72, 72), content)
        document.save(source)
        document.close()
        return source

    def test_pdf_word_text_images_and_page_selection(self):
        from docx import Document
        source = self.pdf()
        output = Path(self.convert(source, "docx")["output"])
        text = "\n".join(p.text for p in Document(output).paragraphs)
        self.assertIn("First page", text)
        self.assertIn("Second page", text)
        result = self.convert(source, "png", pages="2")
        self.assertEqual(len(result["outputs"]), 1)
        result = self.convert(source, "txt")
        self.assertIn("Second page", Path(result["output"]).read_text(encoding="utf-8"))

    def test_scanned_pdf_requires_ocr(self):
        import pymupdf
        from PIL import Image
        image = self.root / "scan.png"
        Image.new("RGB", (40, 40), "white").save(image)
        source = self.root / "scan.pdf"
        document = pymupdf.open()
        page = document.new_page()
        page.insert_image(page.rect, filename=str(image))
        document.save(source)
        document.close()
        with self.assertRaisesRegex(ValueError, "OCR"):
            self.convert(source, "docx")

    def test_pdf_table_to_xlsx(self):
        import pymupdf
        import openpyxl
        source = self.root / "table.pdf"
        document = pymupdf.open()
        page = document.new_page()
        for x in [50, 200, 350]:
            page.draw_line((x, 100), (x, 220))
        for y in [100, 140, 180, 220]:
            page.draw_line((50, y), (350, y))
        for row, values in enumerate([["code", "value"], ["0001", "25.50"], ["0002", "31.20"]]):
            for col, value in enumerate(values):
                page.insert_text((60 + col * 150, 125 + row * 40), value)
        document.save(source)
        document.close()
        output = self.convert(source, "xlsx")
        book = openpyxl.load_workbook(output["output"])
        self.assertEqual(book.active["A2"].value, "0001")
        self.assertEqual(book.active["B2"].value, "25.50")
        book.close()

    def test_pdf_image_pptx_is_explicit(self):
        from pptx import Presentation
        source = self.pdf()
        with self.assertRaisesRegex(ValueError, "image slides"):
            self.convert(source, "pptx")
        output = self.convert(source, "pptx", allow_layout_loss=True)
        self.assertEqual(len(Presentation(output["output"]).slides), 2)

    def test_new_image_formats_heic_animation_and_icon(self):
        from PIL import Image
        import pillow_heif
        pillow_heif.register_heif_opener()
        source = self.root / "source.png"
        Image.new("RGB", (64, 64), "red").save(source)
        for target in ["bmp", "tiff", "gif", "ico", "avif"]:
            if target in self.available["image_targets"]:
                with self.subTest(target=target):
                    self.convert(source, target)
        heic = self.root / "phone.heic"
        Image.open(source).save(heic, format="HEIF")
        self.convert(heic, "jpg")
        animated = self.root / "motion.gif"
        Image.new("RGB", (32, 32), "red").save(animated, save_all=True,
            append_images=[Image.new("RGB", (32, 32), "blue")], duration=[100, 200], loop=0)
        result = self.convert(animated, "webp", all_frames=True)
        with Image.open(result["output"]) as image:
            self.assertEqual(image.n_frames, 2)

    def test_image_merge_preserves_order(self):
        from PIL import Image
        import pymupdf
        sources = []
        for name, color in [("second", "red"), ("first", "blue")]:
            source = self.root / (name + ".png")
            Image.new("RGB", (16, 16), color).save(source)
            sources.append(source)
        args = argparse.Namespace(input=list(map(str, sources)), format="pdf", output_dir=None,
                                  background=None, all_frames=False)
        result = converter.merge_images(args, self.available)
        with pymupdf.open(result["output"]) as document:
            self.assertEqual(len(document), 2)
            pixmap = document[0].get_pixmap()
            self.assertGreater(pixmap.pixel(8, 8)[0], pixmap.pixel(8, 8)[2])

    def test_subtitle_timing(self):
        import pysubs2
        source = self.root / "captions.srt"
        source.write_text("1\n00:00:01,000 --> 00:00:02,500\n中文 caption\n", encoding="utf-8")
        for target in ["vtt", "ass", "ssa"]:
            result = self.convert(source, target)
            subtitles = pysubs2.load(result["output"], encoding="utf-8")
            self.assertEqual(subtitles[0].start, 1000)
            self.assertEqual(subtitles[0].end, 2500)
            self.assertEqual(subtitles[0].text, "中文 caption")

    def test_markdown_docx_epub(self):
        from docx import Document
        source = self.root / "notes.md"
        source.write_text("# 标题\n\nSome original text.\n", encoding="utf-8")
        docx = self.convert(source, "docx")
        self.assertIn("Some original text", "\n".join(p.text for p in Document(docx["output"]).paragraphs))
        epub = Path(self.convert(source, "epub")["output"])
        self.convert(epub, "html")
        self.convert(epub, "pdf")

    def test_svg_and_font_containers(self):
        from PIL import Image
        source = self.root / "drawing.svg"
        source.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="40" height="30"><rect width="40" height="30" fill="red"/></svg>', encoding="utf-8")
        result = self.convert(source, "png", dpi=72)
        with Image.open(result["output"]) as output:
            self.assertEqual(output.size, (40, 30))
        font = Path("C:/Windows/Fonts/arial.ttf")
        if font.exists():
            woff = Path(self.convert(font, "woff2", output_dir=str(self.root))["output"])
            result = self.convert(woff, "ttf")
            self.assertEqual(result["validation"], "font_glyphs_reopened")

    def test_font_targets_respect_outline_type(self):
        source = self.root / "outline.otf"
        source.write_bytes(b"OTTO" + b"\0" * 12)
        targets = converter.targets_for(source, self.available)
        self.assertIn("otf", targets)
        self.assertNotIn("ttf", targets)

    def test_macro_workbook_requires_explicit_loss_choice(self):
        import openpyxl
        source = self.root / "macro.xlsm"
        workbook = openpyxl.Workbook()
        workbook.save(source)
        with self.assertRaisesRegex(ValueError, "removes macros"):
            self.convert(source, "xlsx")

    def test_archives_repack_and_traversal_rejection(self):
        import tarfile
        source = self.root / "archive.zip"
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("folder/中文.txt", "original text")
        tar = Path(self.convert(source, "tar.gz")["output"])
        restored = Path(self.convert(tar, "zip")["output"])
        with zipfile.ZipFile(restored) as archive:
            self.assertEqual(archive.read("folder/中文.txt"), b"original text")
        seven = Path(self.convert(source, "7z")["output"])
        rezipped = Path(self.convert(seven, "zip")["output"])
        with zipfile.ZipFile(rezipped) as archive:
            self.assertEqual(archive.read("folder/中文.txt"), b"original text")
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("../outside.txt", "bad")
        with self.assertRaisesRegex(ValueError, "unsafe path"):
            self.convert(source, "zip")
        self.assertFalse((self.root.parent / "outside.txt").exists())

    def test_expanded_media_targets(self):
        ffmpeg = converter.ffmpeg_path()
        source = self.root / "clip.mp4"
        converter.run([ffmpeg, "-v", "error", "-y", "-f", "lavfi", "-i", "color=c=blue:s=32x24:d=0.4",
                       "-f", "lavfi", "-i", "sine=frequency=440:duration=0.4", "-shortest",
                       "-c:v", "libx264", "-c:a", "aac", str(source)])
        for target in self.available["media_targets"]:
            with self.subTest(target=target):
                self.convert(source, target, allow_media_loss=True)
        self.convert(source, "gif", duration=0.2, start=0, fps=5)
        self.convert(source, "png", start=0)

    def test_office_native_pdf_and_legacy_exports(self):
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches
        import openpyxl
        import pymupdf
        docx = self.root / "report.docx"
        document = Document()
        document.add_heading("Report title", 0)
        document.add_paragraph("Original report content.")
        document.save(docx)
        sources = [docx]
        xlsx = self.root / "workbook.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active.append(["Name", "Amount"])
        workbook.active.append(["Test", 25])
        workbook.active.append(["Total", "=SUM(B2:B2)"])
        workbook.save(xlsx)
        sources.append(xlsx)
        pptx = self.root / "slides.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Original slide content"
        deck.save(pptx)
        sources.append(pptx)
        for source in sources:
            with self.subTest(source=source.suffix):
                result = self.convert(source, "pdf")
                with pymupdf.open(result["output"]) as document:
                    self.assertGreater(len(document), 0)
                    self.assertTrue(any(page.get_text().strip() for page in document))
        for source, target in [(docx, "odt"), (docx, "rtf"), (xlsx, "xls"), (pptx, "odp")]:
            with self.subTest(target=target):
                result = self.convert(source, target)
                self.assertTrue(Path(result["output"]).exists())


if __name__ == "__main__":
    unittest.main()
