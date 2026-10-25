import argparse
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "convert.py"
spec = importlib.util.spec_from_file_location("cross_converter", SCRIPT)
converter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(converter)


class CrossTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.available = converter.capabilities()

    def setUp(self):
        self.workspace = tempfile.TemporaryDirectory(prefix="converter-cross-")
        self.root = Path(self.workspace.name)

    def tearDown(self):
        self.assertEqual(list(self.root.glob(".file-converter-*")), [])
        self.workspace.cleanup()

    def convert(self, source, target, **options):
        original = hashlib.sha256(source.read_bytes()).hexdigest()
        try:
            return converter.convert(source, argparse.Namespace(format=target, **options), self.available)
        finally:
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original)

    def ocr_image(self):
        from PIL import Image, ImageDraw, ImageFont
        source = self.root / "scan.png"
        image = Image.new("RGB", (1400, 600), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 48)
        draw.text((50, 50), "Order 000123", fill="black", font=font)
        draw.text((50, 150), "Total 128.50", fill="black", font=font)
        chinese_font = Path("C:/Windows/Fonts/msyh.ttc")
        if chinese_font.exists():
            draw.text((50, 250), "文件转换测试", fill="black", font=ImageFont.truetype(str(chinese_font), 48))
        image.save(source)
        return source

    def test_image_ocr_text_docx_and_searchable_pdf(self):
        from docx import Document
        import pymupdf
        source = self.ocr_image()
        text = self.convert(source, "txt")
        self.assertIn("128.50", Path(text["output"]).read_text(encoding="utf-8"))
        if Path("C:/Windows/Fonts/msyh.ttc").exists():
            self.assertIn("文件转换", Path(text["output"]).read_text(encoding="utf-8"))
        document = self.convert(source, "docx")
        self.assertIn("Order", "\n".join(p.text for p in Document(document["output"]).paragraphs))
        searchable = self.convert(source, "pdf", ocr=True)
        with pymupdf.open(searchable["output"]) as pdf:
            self.assertIn("128.50", "".join(page.get_text() for page in pdf))

    def test_scanned_pdf_ocr_and_editable_text(self):
        import pymupdf
        from docx import Document
        image = self.ocr_image()
        source = self.root / "scanned.pdf"
        document = pymupdf.open()
        page = document.new_page(width=700, height=300)
        page.insert_image(page.rect, filename=str(image))
        document.save(source)
        document.close()
        result = self.convert(source, "docx", ocr=True)
        self.assertIn("128.50", "\n".join(p.text for p in Document(result["output"]).paragraphs))
        result = self.convert(source, "pdf", ocr=True)
        with pymupdf.open(result["output"]) as pdf, pymupdf.open(source) as original:
            self.assertTrue(pdf[0].get_text().strip())
            self.assertEqual(pdf[0].get_pixmap().samples, original[0].get_pixmap().samples)

    def test_table_screenshot_ocr(self):
        from PIL import Image, ImageDraw, ImageFont
        import openpyxl
        source = self.root / "table.png"
        image = Image.new("RGB", (1300, 500), "white")
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 42)
        for row, values in enumerate([["Code", "Quantity", "Amount"], ["000123", "2", "128.50"], ["000124", "3", "240.00"]]):
            for col, value in enumerate(values):
                draw.text((50 + col * 420, 50 + row * 130), value, fill="black", font=font)
        image.save(source)
        result = self.convert(source, "xlsx")
        workbook = openpyxl.load_workbook(result["output"])
        values = [str(cell.value) for row in workbook.active for cell in row if cell.value is not None]
        self.assertIn("128.50", values)
        self.assertIn("000123", values)
        workbook.close()

    def test_bitmap_tracing_creates_paths(self):
        from PIL import Image, ImageDraw
        source = self.root / "logo.png"
        image = Image.new("RGB", (100, 100), "white")
        ImageDraw.Draw(image).ellipse((20, 20, 80, 80), fill="black")
        image.save(source)
        with self.assertRaisesRegex(ValueError, "monochrome"):
            self.convert(source, "svg")
        result = self.convert(source, "svg", allow_layout_loss=True)
        content = Path(result["output"]).read_text(encoding="utf-8")
        self.assertIn("<path", content)
        self.assertNotIn("<image", content)

    def test_embedded_image_extraction(self):
        from PIL import Image
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches
        import pymupdf
        picture = self.root / "picture.png"
        Image.new("RGB", (40, 30), "red").save(picture)
        word = self.root / "document.docx"
        document = Document()
        document.add_picture(str(picture))
        document.save(word)
        self.assertTrue(Path(self.convert(word, "images")["output"]).suffix == ".png")
        slides = self.root / "slides.pptx"
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(str(picture), Inches(1), Inches(1))
        deck.save(slides)
        self.convert(slides, "images")
        source = self.root / "image.pdf"
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_image(page.rect, filename=str(picture))
        pdf.save(source)
        pdf.close()
        self.convert(source, "images")

    def test_frame_sequence_and_image_video(self):
        from PIL import Image
        first, second = self.root / "a.png", self.root / "b.png"
        Image.new("RGB", (32, 24), "red").save(first)
        Image.new("RGB", (40, 32), "blue").save(second)
        args = argparse.Namespace(input=[str(first), str(second)], format="mp4", fps=2,
                                  output_dir=None, background=None, all_frames=False)
        result = converter.merge_images(args, self.available)
        video = Path(result["output"])
        frames = self.convert(video, "frames", duration=1, fps=2)
        self.assertEqual(len(frames["outputs"]), 2)
        self.assertTrue(all(Path(path).suffix == ".png" for path in frames["outputs"]))
        animation = self.root / "animation.gif"
        Image.open(first).save(animation, save_all=True, append_images=[Image.new("RGB", (32, 24), "blue")], duration=200, loop=0)
        self.convert(animation, "mp4")

    def test_asr_uses_local_model_and_preserves_source(self):
        model = Path(os.environ.get("FILE_CONVERTER_ASR_MODEL",
                                    str(Path(__file__).resolve().parents[1] / ".cache/models/tiny.en")))
        if not (model / "model.bin").is_file():
            self.skipTest("Local ASR model is not bundled with skill")
        import win32com.client
        source = self.root / "speech.wav"
        voice = win32com.client.Dispatch("SAPI.SpVoice")
        for token in voice.GetVoices():
            if "409" in token.GetAttribute("Language").split(";"):
                voice.Voice = token
                break
        stream = win32com.client.Dispatch("SAPI.SpFileStream")
        stream.Format.Type = 22
        stream.Open(str(source), 3, False)
        try:
            voice.AudioOutputStream = stream
            voice.Speak("This is a file conversion test. The total amount is one hundred twenty eight dollars.", 0)
        finally:
            stream.Close()
        result = self.convert(source, "txt", asr_model_path=str(model), language="en")
        self.assertIn("conversion", Path(result["output"]).read_text(encoding="utf-8").lower())
        result = self.convert(source, "srt", asr_model_path=str(model), language="en")
        import pysubs2
        self.assertGreater(len(pysubs2.load(result["output"], encoding="utf-8")), 0)
        with self.assertRaisesRegex(ValueError, "English-only"):
            self.convert(source, "txt", asr_model_path=str(model), language="zh")

    def test_office_pages_and_markdown_pdf(self):
        from PIL import Image
        from docx import Document
        source = self.root / "report.docx"
        document = Document()
        document.add_paragraph("First page")
        document.add_page_break()
        document.add_paragraph("Second page")
        document.save(source)
        result = self.convert(source, "png", pages="2")
        self.assertEqual(len(result["outputs"]), 1)
        with Image.open(result["output"]) as picture:
            self.assertGreater(picture.width, 0)
        markdown = self.root / "notes.md"
        markdown.write_text("# Title\n\nPDF content.\n", encoding="utf-8")
        result = self.convert(markdown, "pdf")
        import pymupdf
        with pymupdf.open(result["output"]) as pdf:
            self.assertIn("PDF content", "".join(page.get_text() for page in pdf))


if __name__ == "__main__":
    unittest.main()
